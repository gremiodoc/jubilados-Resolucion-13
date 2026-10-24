#!/usr/bin/env python3
"""
Genera assets/js/data.js a partir del consolidado de jubilados en Excel.

Uso:
    python herramientas/generar_datos.py --xlsx "lista_jubilados_resolucion_13-v2.xlsx"
    python herramientas/generar_datos.py --xlsx "lista_jubilados_resolucion_13-v2.xlsx" --out assets/js/data.js

El script detecta automáticamente la hoja y las columnas de cédula y categoría.
Si existe una hoja llamada "Listado Completo", la usa por defecto.

El archivo de salida NO contiene cédulas en texto plano: guarda huellas de
64 bits (dos variantes FNV-1a) agrupadas por tipo de personal. La función
hash_cedula() de este script debe coincidir exactamente con hashCedula()
de assets/js/app.js.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("Falta openpyxl. Instálalo con: pip install openpyxl")


TYPE_NAMES = {
    "D": "Personal Docente",
    "A": "Personal Administrativo",
    "O": "Personal Obrero",
}

PREFERRED_SHEETS = ("Listado Completo", "Listado completo", "Hoja1")


def only_digits(value: object) -> str:
    return re.sub(r"\D", "", str(value or ""))


def normalize_text(value: object) -> str:
    """Minúsculas sin acentos, para comparar encabezados y categorías."""
    decomposed = unicodedata.normalize("NFKD", str(value or ""))
    plain = "".join(char for char in decomposed if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", plain).strip().lower()


def classify(category: str) -> str | None:
    text = normalize_text(category)
    if "docente" in text:
        return "D"
    if "administrativo" in text:
        return "A"
    if "obrero" in text:
        return "O"
    return None


def hash_cedula(digits: str) -> str:
    """Huella de 64 bits: FNV-1a + segunda variante, con avalancha final."""
    h1 = 0x811C9DC5
    h2 = 0x9747B28C
    for char in digits:
        code = ord(char)
        h1 = ((h1 ^ code) * 0x01000193) & 0xFFFFFFFF
        h2 = ((h2 ^ code) * 0x5BD1E995) & 0xFFFFFFFF
        h2 ^= h2 >> 15
        h2 &= 0xFFFFFFFF
    h1 ^= h1 >> 16
    h1 = (h1 * 0x85EBCA6B) & 0xFFFFFFFF
    h1 ^= h1 >> 13
    h2 ^= h2 >> 13
    h2 = (h2 * 0xC2B2AE35) & 0xFFFFFFFF
    h2 ^= h2 >> 16
    return f"{h1:08x}{h2:08x}"


def js_array(items: list[str], per_line: int = 10) -> str:
    lines = []
    for index in range(0, len(items), per_line):
        chunk = ", ".join(json.dumps(item) for item in items[index:index + per_line])
        comma = "," if index + per_line < len(items) else ""
        lines.append("    " + chunk + comma)
    return "\n".join(lines)


def choose_sheet(workbook, sheet_name: str | None):
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(
                f"No existe la hoja {sheet_name!r}. Hojas disponibles: "
                + ", ".join(workbook.sheetnames)
            )
        return workbook[sheet_name]
    for candidate in PREFERRED_SHEETS:
        if candidate in workbook.sheetnames:
            return workbook[candidate]
    return workbook.active


def find_columns(rows: list[tuple]) -> tuple[int, int, int]:
    """Devuelve (fila_encabezado, columna_categoria, columna_cedula), base 0."""
    for row_index, row in enumerate(rows[:25]):
        category_index = None
        category_score = 0
        cedula_index = None
        for column_index, value in enumerate(row):
            text = normalize_text(value)
            score = 0
            if text == "categoria":
                score = 4
            elif text.startswith("categoria"):
                score = 3
            elif "categoria" in text:
                score = 2
            elif text == "personal":
                score = 2
            elif "personal" in text:
                score = 1
            if score > category_score:
                category_score = score
                category_index = column_index
            if "cedula" in text or "dula" in text:
                if cedula_index is None:
                    cedula_index = column_index
        if category_index is not None and cedula_index is not None:
            return row_index, category_index, cedula_index

    # Respaldo para archivos con encabezados dañados o con otro formato:
    # se eligen las columnas con más valores de categoría y más cédulas V-/E-.
    column_count = max((len(row) for row in rows[:200]), default=0)
    category_scores = [0] * column_count
    cedula_scores = [0] * column_count
    for row in rows[:200]:
        for column_index, value in enumerate(row):
            if classify(value):
                category_scores[column_index] += 1
            text = str(value or "").strip()
            if re.match(r"^[VEve]\s*[-–]?\s*\d", text):
                cedula_scores[column_index] += 1

    if not category_scores or max(category_scores) == 0 or max(cedula_scores) == 0:
        raise ValueError(
            "No se encontraron las columnas de Cédula y Categoría. "
            "Se revisaron encabezados y contenido de las primeras filas."
        )

    category_index = category_scores.index(max(category_scores))
    cedula_index = cedula_scores.index(max(cedula_scores))
    print("Columnas detectadas por contenido (encabezado no reconocido).")
    return 0, category_index, cedula_index


def read_records(
    xlsx: Path, sheet_name: str | None
) -> tuple[dict[str, list[str]], Counter, str, int]:
    workbook = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    try:
        sheet = choose_sheet(workbook, sheet_name)
        rows = list(sheet.iter_rows(values_only=True))
        header_index, category_index, cedula_index = find_columns(rows)

        records: dict[str, list[str]] = {"D": [], "A": [], "O": []}
        seen: dict[str, str] = {}
        skipped_total = 0

        for row in rows[header_index + 1:]:
            if max(category_index, cedula_index) >= len(row):
                continue
            category = row[category_index]
            cedula = row[cedula_index]
            if category is None or cedula is None:
                continue
            category_text = normalize_text(category)
            if not category_text:
                continue
            if "todas" in category_text or "total" in category_text:
                skipped_total += 1
                continue
            code = classify(category_text)
            if code is None:
                # Títulos, encabezados repetidos y filas ajenas al consolidado.
                continue
            digits = only_digits(cedula)
            if not digits:
                continue
            if len(digits) < 5 or len(digits) > 9:
                raise ValueError(f"Cédula con longitud inesperada: {cedula!r}")
            if digits in seen:
                raise ValueError(
                    f"Cédula duplicada: {cedula!r} ya aparece como {seen[digits]}"
                )
            seen[digits] = code
            records[code].append(digits)

        if not any(records.values()):
            raise ValueError(
                "No se encontraron registros. Verifica que la hoja tenga las columnas "
                "de Cédula y Categoría, y que las categorías incluyan Docente, "
                "Administrativo u Obrero."
            )

        counts = Counter({code: len(values) for code, values in records.items()})
        if skipped_total:
            print(f"Filas de total ignoradas: {skipped_total}")
        print(
            f"Hoja usada: {sheet.title} | encabezado en fila {header_index + 1} | "
            f"categoría columna {category_index + 1} | cédula columna {cedula_index + 1}"
        )
        return records, counts, sheet.title, header_index
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", required=True, help="Ruta del archivo Excel de origen")
    parser.add_argument(
        "--out",
        default="assets/js/data.js",
        help="Ruta de salida para data.js (por defecto: assets/js/data.js)",
    )
    parser.add_argument("--sheet", default=None, help="Nombre de la hoja (opcional)")
    parser.add_argument(
        "--version",
        type=int,
        default=2,
        help="Versión interna de la base (por defecto: 2)",
    )
    args = parser.parse_args()

    xlsx = Path(args.xlsx).expanduser().resolve()
    out = Path(args.out).expanduser()
    if not out.is_absolute():
        out = (Path.cwd() / out).resolve()

    if not xlsx.is_file():
        sys.exit(f"No existe el archivo: {xlsx}")

    records, counts, sheet_name, _ = read_records(xlsx, args.sheet)

    hashed = {
        code: [hash_cedula(digits) for digits in sorted(records[code], key=int)]
        for code in ("D", "A", "O")
    }
    all_hashes = [value for values in hashed.values() for value in values]
    if len(set(all_hashes)) != len(all_hashes):
        sys.exit("Se produjo una colisión de huellas. Cambia el algoritmo antes de continuar.")

    total = sum(counts.values())
    data_js = f"""/**
 * Listado de Jubilados — Resolución N.º 013 · Agosto 2026
 * Ministerio del Poder Popular para la Educación · República Bolivariana de Venezuela
 *
 * Base de consulta generada a partir del archivo oficial «{xlsx.name}».
 * Hoja utilizada: {sheet_name}.
 * Por privacidad, este archivo NO contiene cédulas en texto plano: guarda
 * índices de 64 bits (FNV-1a doble) agrupados por tipo de personal.
 *
 * D = Docente ({counts['D']:,} registros)
 * A = Administrativo ({counts['A']:,} registros)
 * O = Obrero ({counts['O']:,} registros)
 * Total = {total:,} registros
 *
 * Este archivo se carga con una etiqueta <script> clásica para funcionar
 * tanto en GitHub Pages como al abrir index.html directamente.
 */
window.JUBILADOS = Object.freeze({{
  resolucion: "013",
  fecha: "13 de agosto de 2026",
  version: {args.version},
  total: {total},
  D: [
{js_array(hashed['D'])}
  ],
  A: [
{js_array(hashed['A'])}
  ],
  O: [
{js_array(hashed['O'])}
  ]
}});
"""

    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(data_js, encoding="utf-8", newline="\n")

    print(f"Archivo generado: {out}")
    print(f"Total: {total:,} registros")
    for code in ("D", "A", "O"):
        print(f"  {code} - {TYPE_NAMES[code]}: {counts[code]:,}")
    print(f"Huellas únicas: {len(set(all_hashes)):,}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
